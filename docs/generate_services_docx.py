"""Generate Word report with code captures for AutoLoc services."""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(r"c:\Users\SOS-HOME\Desktop\autoloc-api")
SERVICE = ROOT / "src" / "main" / "java" / "tn" / "esprit" / "autoloc" / "autolocapi" / "service"
OUT_DIR = ROOT / "docs" / "captures_services"
OUT_DOCX = ROOT / "docs" / "Services_et_Implementations_AutoLoc.docx"

# One screenshot per file, like the associations document.
FILES = [
    ("IAgenceService", "IAgenceService.java", "Interface du service Agence"),
    ("IAgenceServiceImp", "IAgenceServiceImp.java", "Implémentation du service Agence"),
    ("IClientService", "IClientService.java", "Interface du service Client"),
    ("IClientServiceImp", "IClientServiceImp.java", "Implémentation du service Client"),
    ("IContratService", "IContratService.java", "Interface du service Contrat"),
    ("IContratServiceImp", "IContratServiceImp.java", "Implémentation du service Contrat"),
    ("IEmployeService", "IEmployeService.java", "Interface du service Employe"),
    ("IEmployeServiceImp", "IEmployeServiceImp.java", "Implémentation du service Employe"),
    ("IEquipementService", "IEquipementService.java", "Interface du service Equipement"),
    ("IEquipementServiceImp", "IEquipementServiceImp.java", "Implémentation du service Equipement"),
    ("IMaintenanceService", "IMaintenanceService.java", "Interface du service Maintenance"),
    ("IMaintenanceServiceImp", "IMaintenanceServiceImp.java", "Implémentation du service Maintenance"),
    ("IPaiementService", "IPaiementService.java", "Interface du service Paiement"),
    ("IPaiementServiceImp", "IPaiementServiceImp.java", "Implémentation du service Paiement"),
    ("IReservationService", "IReservationService.java", "Interface du service Reservation"),
    ("IReservationServiceImp", "IReservationServiceImp.java", "Implémentation du service Reservation"),
    ("IVehiculeService", "IVehiculeService.java", "Interface du service Vehicule"),
    ("IVehiculeServiceImp", "IVehiculeServiceImp.java", "Implémentation du service Vehicule"),
]

BG = (30, 30, 30)
LINE_NUM_BG = (37, 37, 38)
LINE_NUM_FG = (133, 133, 133)
FG = (212, 212, 212)
KEYWORD = (86, 156, 214)
ANNOTATION = (78, 201, 176)
STRING = (206, 145, 120)
TYPE = (78, 201, 176)
COMMENT = (106, 153, 85)
TITLE_BAR = (45, 45, 48)
TITLE_FG = (204, 204, 204)

JAVA_KEYWORDS = {
    "package", "import", "public", "class", "interface", "implements",
    "private", "protected", "new", "return", "static", "final", "void",
    "boolean", "true", "false", "null", "this", "super", "extends",
}


def get_font(size: int, bold: bool = False):
    candidates = [
        r"C:\Windows\Fonts\consola.ttf",
        r"C:\Windows\Fonts\CascadiaMono.ttf",
        r"C:\Windows\Fonts\cour.ttf",
        r"C:\Windows\Fonts\lucon.ttf",
    ]
    bold_candidates = [
        r"C:\Windows\Fonts\consolab.ttf",
        r"C:\Windows\Fonts\courbd.ttf",
    ]
    paths = bold_candidates if bold else candidates
    for path in paths + candidates:
        if Path(path).exists():
            return ImageFont.truetype(path, size)
    return ImageFont.load_default()


def tokenize_line(line: str):
    tokens = []
    i = 0
    n = len(line)
    while i < n:
        if line[i] == '"':
            j = i + 1
            while j < n and line[j] != '"':
                if line[j] == "\\":
                    j += 2
                    continue
                j += 1
            j = min(j + 1, n)
            tokens.append((STRING, line[i:j]))
            i = j
            continue
        if line.startswith("//", i):
            tokens.append((COMMENT, line[i:]))
            break
        if line[i] == "@":
            j = i + 1
            while j < n and (line[j].isalnum() or line[j] == "_"):
                j += 1
            tokens.append((ANNOTATION, line[i:j]))
            i = j
            continue
        if line[i].isalpha() or line[i] == "_":
            j = i + 1
            while j < n and (line[j].isalnum() or line[j] == "_"):
                j += 1
            word = line[i:j]
            if word in JAVA_KEYWORDS:
                tokens.append((KEYWORD, word))
            elif word[:1].isupper():
                tokens.append((TYPE, word))
            else:
                tokens.append((FG, word))
            i = j
            continue
        tokens.append((FG, line[i]))
        i += 1
    return tokens


def render_code_image(filename: str, code: str, out_path: Path):
    font = get_font(15)
    title_font = get_font(14, bold=True)
    lines = code.replace("\r\n", "\n").split("\n")
    if lines and lines[-1] == "":
        lines = lines[:-1]

    pad_x = 16
    pad_y = 12
    title_h = 36
    line_h = 22
    gutter_w = 52
    tmp = Image.new("RGB", (10, 10))
    td = ImageDraw.Draw(tmp)
    max_w = 0
    for line in lines:
        max_w = max(max_w, int(td.textlength(line.replace("\t", "    "), font=font)))

    width = gutter_w + pad_x * 2 + max(max_w + 40, 900)
    height = title_h + pad_y * 2 + line_h * len(lines) + 8

    img = Image.new("RGB", (width, height), BG)
    draw = ImageDraw.Draw(img)

    draw.rectangle([0, 0, width, title_h], fill=TITLE_BAR)
    draw.ellipse([12, 12, 22, 22], fill=(255, 95, 86))
    draw.ellipse([28, 12, 38, 22], fill=(255, 189, 46))
    draw.ellipse([44, 12, 54, 22], fill=(39, 201, 63))
    draw.text((70, 9), f"{filename}  —  AutoLoc", font=title_font, fill=TITLE_FG)

    draw.rectangle([0, title_h, gutter_w, height], fill=LINE_NUM_BG)

    y = title_h + pad_y
    for idx, line in enumerate(lines, start=1):
        line = line.replace("\t", "    ")
        num = str(idx)
        nw = int(draw.textlength(num, font=font))
        draw.text((gutter_w - 10 - nw, y), num, font=font, fill=LINE_NUM_FG)
        x = gutter_w + pad_x
        for color, token in tokenize_line(line):
            draw.text((x, y), token, font=font, fill=color)
            x += int(draw.textlength(token, font=font))
        y += line_h

    out_path.parent.mkdir(parents=True, exist_ok=True)
    img.save(out_path, "PNG")
    return out_path


def set_run_font(run, name="Calibri", size=11, bold=False):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold


def add_page_break(doc):
    p = doc.add_paragraph()
    run = p.add_run()
    run.add_break(WD_BREAK.PAGE)


def main():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    doc = Document()

    section = doc.sections[0]
    section.top_margin = Cm(1.5)
    section.bottom_margin = Cm(1.5)
    section.left_margin = Cm(1.5)
    section.right_margin = Cm(1.5)

    title = doc.add_heading("AutoLoc — Services et implémentations", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(
        "Captures du code de chaque fichier de la couche service : "
        "interface (IXxxService) puis implémentation (IXxxServiceImp)."
    )
    set_run_font(run, size=12)

    doc.add_heading("Récapitulatif", level=1)
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    hdr[0].text = "Fichier"
    hdr[1].text = "Type"
    hdr[2].text = "Description"
    for name, fname, desc in FILES:
        row = table.add_row().cells
        row[0].text = fname
        row[1].text = "Interface" if fname.endswith("Service.java") else "Implémentation"
        row[2].text = desc

    add_page_break(doc)
    doc.add_heading("Captures du code par fichier", level=1)

    for i, (name, fname, desc) in enumerate(FILES):
        path = SERVICE / fname
        code = path.read_text(encoding="utf-8")
        img_path = OUT_DIR / f"{name}.png"
        render_code_image(fname, code, img_path)

        kind = "Interface" if fname.endswith("Service.java") else "Implémentation"
        doc.add_heading(f"{kind} : {name}", level=2)

        info = doc.add_paragraph()
        r1 = info.add_run("Fichier : ")
        set_run_font(r1, bold=True, size=11)
        r2 = info.add_run(fname)
        set_run_font(r2, size=11)

        info2 = doc.add_paragraph()
        r3 = info2.add_run("Description : ")
        set_run_font(r3, bold=True, size=11)
        r4 = info2.add_run(desc)
        set_run_font(r4, size=11)

        pic = doc.add_picture(str(img_path), width=Cm(17.5))
        last = doc.paragraphs[-1]
        last.alignment = WD_ALIGN_PARAGRAPH.CENTER

        if i < len(FILES) - 1:
            add_page_break(doc)

    OUT_DOCX.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT_DOCX)
    print(f"OK: {OUT_DOCX}")
    print(f"Captures: {OUT_DIR}")
    print(f"Fichiers: {len(FILES)}")


if __name__ == "__main__":
    main()
